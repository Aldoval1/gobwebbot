import os
import requests
from datetime import datetime, timedelta
from flask import render_template, flash, redirect, url_for, request, current_app, jsonify, make_response, session
from app import db
from app.forms import (
    LoginForm, RegistrationForm, OfficialLoginForm, OfficialRegistrationForm,
    SearchUserForm, CriminalRecordForm, TrafficFineForm, CommentForm,
    AppointmentForm, CreateLeaderForm, EditCitizenForm, EditCitizenPhotoForm,
    BusinessLicenseForm, UserPhotoForm, ChangePasswordForm
)
from app.models import (
    User, Comment, TrafficFine, License, CriminalRecord,
    CriminalRecordSubjectPhoto, CriminalRecordEvidencePhoto,
    Appointment, Business, BusinessFine, Document as DocModel
)
from sqlalchemy import func
from flask_login import current_user, login_user, logout_user, login_required
from flask import Blueprint
from werkzeug.utils import secure_filename
from fpdf import FPDF
from docx import Document
from pypdf import PdfReader
import io
from flask import send_file

bp = Blueprint('main', __name__)

# --- CONFIGURACIÓN DISCORD OAUTH2 ---
DISCORD_CLIENT_ID = os.environ.get('DISCORD_CLIENT_ID')
DISCORD_CLIENT_SECRET = os.environ.get('DISCORD_CLIENT_SECRET')
BASE_URL = os.environ.get('WEB_APP_URL', 'http://127.0.0.1:5000')
DISCORD_REDIRECT_URI = f"{BASE_URL}/callback"
DISCORD_API_ENDPOINT = 'https://discord.com/api/v10'

# --- HELPER FUNCTIONS ---

def notify_discord_bot(user, message):
    if not user.discord_id:
        return
    
    bot_url = os.environ.get('BOT_URL')
    if not bot_url:
        print("ADVERTENCIA: Variable 'BOT_URL' no configurada en la Web.")
        return

    try:
        payload = {
            'discord_id': user.discord_id,
            'message': message
        }
        requests.post(f"{bot_url}/notify", json=payload, timeout=2)
    except Exception as e:
        print(f"Error enviando notificación a Discord: {e}")

def _perform_user_deletion(user):
    # 1. Nullify author_id in related records to avoid deletion or integrity errors
    TrafficFine.query.filter_by(author_id=user.id).update({TrafficFine.author_id: None})
    CriminalRecord.query.filter_by(author_id=user.id).update({CriminalRecord.author_id: None})
    Comment.query.filter_by(author_id=user.id).update({Comment.author_id: None})
    DocModel.query.filter_by(uploader_id=user.id).update({DocModel.uploader_id: None})

    # 2. Delete user (Cascade will handle owned records)
    db.session.delete(user)

@bp.route('/official/toggle_duty', methods=['POST'])
@login_required
def official_toggle_duty():
    allowed_departments = ['SABES', 'Gobierno', 'Ejecutivo', 'Legislativo', 'Judicial']
    if current_user.department not in allowed_departments:
        flash('Acceso denegado.')
        return redirect(url_for('main.official_dashboard'))

    current_user.on_duty = not current_user.on_duty
    db.session.commit()

    status_msg = "EN SERVICIO" if current_user.on_duty else "FUERA DE SERVICIO"
    flash(f'Estado actualizado: {status_msg}')

    if current_user.on_duty:
        # Notificar a usuarios suscritos
        subscribed_users = User.query.filter(User.discord_id.isnot(None), User.receive_notifications == True).all()

        # En un entorno real, esto debería ser una tarea en segundo plano (Celery/Redis)
        # para no bloquear la respuesta HTTP si hay muchos usuarios.
        # Aquí simulamos enviando solo a los primeros 50 para evitar timeouts en este MVP.
        count = 0
        link = url_for('main.settings_notifications', _external=True)
        message = (
            f"👮 **{current_user.department}**\n"
            f"El funcionario **{current_user.first_name} {current_user.last_name}** está ahora en servicio.\n\n"
            f"Gestionar notificaciones: {link}"
        )

        for user in subscribed_users[:50]:
            notify_discord_bot(user, message)
            count += 1

        print(f"Notificación de servicio enviada a {count} usuarios via Discord.")

    return redirect(url_for('main.official_dashboard'))

@bp.route('/settings/notifications', methods=['GET', 'POST'])
@login_required
def settings_notifications():
    if request.method == 'POST':
        # Simple toggle via form submit or check
        enable = request.form.get('receive_notifications') == 'on'
        current_user.receive_notifications = enable
        db.session.commit()
        flash('Preferencias de notificación actualizadas.')
        return redirect(url_for('main.index'))

    return render_template('settings.html')

# --- API ROUTES FOR DISCORD BOT ---

@bp.route('/api/check_citizen/<dni>', methods=['GET'])
def check_citizen_api(dni):
    user = User.query.filter_by(dni=dni).first()
    if user:
        return jsonify({
            'found': True,
            'first_name': user.first_name,
            'last_name': user.last_name
        })
    return jsonify({'found': False}), 404

@bp.route('/api/link_discord', methods=['POST'])
def link_discord_api():
    data = request.get_json()
    dni = data.get('dni')
    discord_id = data.get('discord_id')
    
    user = User.query.filter_by(dni=dni).first()
    if user:
        user.discord_id = str(discord_id)
        db.session.commit()
        return jsonify({'success': True})
    
    return jsonify({'success': False, 'message': 'Usuario no encontrado'}), 404

@bp.route('/api/search_users', methods=['GET'])
@login_required
def api_search_users():
    query = request.args.get('q', '')
    if not query or len(query) < 2:
        return jsonify([])

    search = f"%{query}%"
    users = User.query.filter(
        (User.first_name.ilike(search)) |
        (User.last_name.ilike(search)) |
        (User.dni.ilike(search))
    ).limit(10).all()

    results = [{'dni': u.dni, 'name': f"{u.first_name} {u.last_name}"} for u in users]
    return jsonify(results)

# --- DISCORD OAUTH2 ROUTES ---

@bp.route('/discord/login')
@login_required
def discord_login():
    if not DISCORD_CLIENT_ID or not DISCORD_CLIENT_SECRET:
        flash('Error: Faltan credenciales de Discord en la configuración (.env).')
        return redirect(url_for('main.citizen_dashboard'))
    
    # Updated scope to include guilds.join
    oauth_url = f"https://discord.com/oauth2/authorize?client_id={DISCORD_CLIENT_ID}&redirect_uri={DISCORD_REDIRECT_URI}&response_type=code&scope=identify%20guilds.join"
    return redirect(oauth_url)

@bp.route('/callback')
@login_required
def discord_callback():
    code = request.args.get('code')
    if not code:
        flash('No se recibió código de autorización de Discord.')
        return redirect(url_for('main.citizen_dashboard'))

    data = {
        'client_id': DISCORD_CLIENT_ID,
        'client_secret': DISCORD_CLIENT_SECRET,
        'grant_type': 'authorization_code',
        'code': code,
        'redirect_uri': DISCORD_REDIRECT_URI
    }
    headers = {'Content-Type': 'application/x-www-form-urlencoded'}
    
    try:
        token_resp = requests.post(f'{DISCORD_API_ENDPOINT}/oauth2/token', data=data, headers=headers)
        token_resp.raise_for_status()
        access_token = token_resp.json().get('access_token')

        user_headers = {'Authorization': f'Bearer {access_token}'}
        user_resp = requests.get(f'{DISCORD_API_ENDPOINT}/users/@me', headers=user_headers)
        user_resp.raise_for_status()

        discord_user_data = user_resp.json()
        discord_id = discord_user_data.get('id')

        current_user.discord_id = discord_id
        db.session.commit()

        # Save access token for guild join step
        session['discord_access_token'] = access_token

        return redirect(url_for('main.discord_select_servers'))

    except requests.exceptions.RequestException as e:
        print(f"Error OAuth Discord: {e}")
        flash('Hubo un error al conectar con Discord. Inténtalo de nuevo.')
        return redirect(url_for('main.citizen_dashboard'))

@bp.route('/discord/select_servers', methods=['GET', 'POST'])
@login_required
def discord_select_servers():
    access_token = session.get('discord_access_token')
    if not access_token:
        # If no token, maybe they already linked before? Just show dashboard
        flash('Sesión de Discord expirada o ya finalizada.')
        return redirect(url_for('main.citizen_dashboard'))

    if request.method == 'POST':
        selected_guilds = request.form.getlist('guilds')

        # Logic to join guilds
        bot_token = current_app.config.get('DISCORD_BOT_TOKEN')

        guild_map = {
            'gobierno': current_app.config.get('GOBIERNO_GUILD_ID'),
            'judicial': current_app.config.get('JUDICIAL_GUILD_ID'),
            'congreso': current_app.config.get('CONGRESO_GUILD_ID')
        }

        # Always try to join Gobierno (mandatory)
        if 'gobierno' not in selected_guilds:
            selected_guilds.append('gobierno')

        success_count = 0

        for key in selected_guilds:
            guild_id = guild_map.get(key)
            if guild_id and bot_token:
                # Add User to Guild using Bot Token + User Access Token
                url = f"https://discord.com/api/v10/guilds/{guild_id}/members/{current_user.discord_id}"
                headers = {
                    "Authorization": f"Bot {bot_token}",
                    "Content-Type": "application/json"
                }
                payload = {"access_token": access_token}

                try:
                    resp = requests.put(url, headers=headers, json=payload, timeout=5)
                    # 201: Joined, 204: Already joined
                    if resp.status_code in [201, 204]:
                        success_count += 1
                    else:
                        print(f"Failed to join guild {key} ({guild_id}): {resp.status_code} {resp.text}")
                except Exception as e:
                    print(f"Error contacting Discord API: {e}")

        # Trigger Bot for Roles & Nicknames
        bot_url = os.environ.get('BOT_URL')
        if bot_url:
            try:
                payload = {
                    'discord_id': current_user.discord_id,
                    'first_name': current_user.first_name,
                    'last_name': current_user.last_name,
                    'guilds': selected_guilds
                }
                requests.post(f"{bot_url}/setup_account", json=payload, timeout=5)
            except Exception as e:
                print(f"Bot setup error: {e}")

        flash(f'¡Configuración completada! Te has unido a los servidores seleccionados.')
        session.pop('discord_access_token', None)
        return redirect(url_for('main.citizen_dashboard'))

    return render_template('select_servers.html')

# --- MAIN ROUTES ---

@bp.route('/')
def index():
    return render_template('landing.html')

@bp.route('/login', methods=['GET', 'POST'])
def login():
    if current_user.is_authenticated:
        if current_user.badge_id:
             return redirect(url_for('main.official_dashboard'))
        return redirect(url_for('main.citizen_dashboard'))

    form = LoginForm()
    if form.validate_on_submit():
        user = User.query.filter_by(dni=form.dni.data, badge_id=None).first()
        if user is None or not user.check_password(form.password.data):
             flash('DNI o contraseña inválidos')
             return redirect(url_for('main.login'))

        login_user(user, remember=form.remember_me.data)
        if user.badge_id:
            return redirect(url_for('main.official_dashboard'))
        return redirect(url_for('main.citizen_dashboard'))

    return render_template('login.html', form=form)

@bp.route('/citizen/dashboard')
@login_required
def citizen_dashboard():
    if current_user.badge_id:
        return redirect(url_for('main.official_dashboard'))
    return render_template('citizen_dashboard.html')

@bp.route('/register', methods=['GET', 'POST'])
def register():
    if current_user.is_authenticated:
        return redirect(url_for('main.index'))

    form = RegistrationForm()
    if form.validate_on_submit():
        user_exist = User.query.filter_by(dni=form.dni.data, badge_id=None).first()
        if user_exist:
            flash('Ese DNI ya está registrado.')
            return redirect(url_for('main.register'))

        selfie_file = form.selfie.data
        dni_photo_file = form.dni_photo.data

        selfie_filename = secure_filename(selfie_file.filename)
        dni_photo_filename = secure_filename(dni_photo_file.filename)

        if not os.path.exists(current_app.config['UPLOAD_FOLDER']):
            os.makedirs(current_app.config['UPLOAD_FOLDER'])

        selfie_path = os.path.join(current_app.config['UPLOAD_FOLDER'], selfie_filename)
        dni_photo_path = os.path.join(current_app.config['UPLOAD_FOLDER'], dni_photo_filename)

        selfie_file.save(selfie_path)
        dni_photo_file.save(dni_photo_path)

        user = User(
            first_name=form.first_name.data,
            last_name=form.last_name.data,
            dni=form.dni.data,
            selfie_filename=selfie_filename,
            dni_photo_filename=dni_photo_filename
        )
        user.set_password(form.password.data)
        db.session.add(user)
        db.session.commit()

        flash('¡Cuenta creada con éxito! Ahora puedes iniciar sesión.')
        return redirect(url_for('main.login'))

    return render_template('register.html', form=form)

@bp.route('/logout')
def logout():
    logout_user()
    return redirect(url_for('main.index'))

# --- CITIZEN FINES ROUTES ---

@bp.route('/my_fines')
@login_required
def my_fines():
    if current_user.badge_id:
        return redirect(url_for('main.official_dashboard'))

    fines = TrafficFine.query.filter_by(user_id=current_user.id, status='Pendiente').all()
    history = TrafficFine.query.filter_by(user_id=current_user.id, status='Pagada').all()

    return render_template('my_fines.html', fines=fines, history=history)

# --- APPOINTMENTS ROUTES ---

@bp.route('/appointments')
@login_required
def appointments():
    if current_user.badge_id:
        return redirect(url_for('main.official_dashboard'))

    officials = User.query.filter_by(department='Gobierno', official_status='Aprobado').all()
    form = AppointmentForm()

    return render_template('appointments.html', officials=officials, form=form)

@bp.route('/appointments/book/<int:official_id>', methods=['POST'])
@login_required
def book_appointment(official_id):
    form = AppointmentForm()
    if form.validate_on_submit():
        official = User.query.get_or_404(official_id)
        if official.department != 'Gobierno':
            flash('Solo puedes solicitar citas con funcionarios del Gobierno.')
            return redirect(url_for('main.appointments'))

        combined_dt = datetime.combine(form.date.data, form.time.data)

        appt = Appointment(
            citizen_id=current_user.id,
            official_id=official.id,
            date=combined_dt,
            reason=form.description.data,
            status='Pending'
        )
        db.session.add(appt)
        db.session.commit()
        
        notify_discord_bot(current_user, f"📅 **Cita Solicitada**\nTu cita con el oficial {official.last_name} ha sido registrada para el {combined_dt}.")
        notify_discord_bot(official, f"📅 **Nueva Cita Recibida**\nEl ciudadano {current_user.first_name} {current_user.last_name} solicita cita para el {combined_dt}.\nMotivo: {form.description.data}")
        
        flash('Cita solicitada con éxito.')
    else:
        flash('Error al solicitar la cita. Revisa los datos.')

    return redirect(url_for('main.appointments'))

# --- MY DOCUMENTS ROUTES ---

@bp.route('/my_documents')
@login_required
def my_documents():
    if current_user.badge_id:
        return redirect(url_for('main.official_dashboard'))

    if current_user.created_at:
        account_age = (datetime.utcnow() - current_user.created_at).days
    else:
        account_age = 0
    
    # Formulario para actualizar foto
    photo_form = UserPhotoForm()

    return render_template('my_documents.html', account_age=account_age, photo_form=photo_form)

@bp.route('/my_documents/update_photo', methods=['POST'])
@login_required
def update_my_photo():
    form = UserPhotoForm()
    if form.validate_on_submit():
        if form.photo.data:
            f = form.photo.data
            filename = secure_filename(f.filename)
            f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], filename))
            current_user.selfie_filename = filename
            db.session.commit()
            flash('Foto de perfil actualizada.')
        else:
            flash('No se seleccionó ninguna foto.')
    else:
        flash('Error al subir la foto. Asegúrate de que sea una imagen válida.')
    
    return redirect(url_for('main.my_documents'))

@bp.route('/my_documents/download_criminal_record')
@login_required
def download_criminal_record():
    records = current_user.criminal_records

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    # Header
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt="Reporte de Antecedentes Penales", ln=True, align='C')
    pdf.ln(10)

    # User Info
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 8, txt=f"Ciudadano: {current_user.first_name} {current_user.last_name}", ln=True)
    pdf.cell(200, 8, txt=f"DNI: {current_user.dni}", ln=True)
    pdf.cell(200, 8, txt=f"Fecha de Emisión: {datetime.utcnow().strftime('%d/%m/%Y %H:%M')}", ln=True)
    pdf.ln(10)

    if not records:
        pdf.cell(200, 10, txt="No se encontraron antecedentes penales.", ln=True)
    else:
        for record in records:
            pdf.set_font("Arial", 'B', 12)
            pdf.cell(200, 8, txt=f"Delito: {record.crime} (CP: {record.penal_code})", ln=True)
            pdf.set_font("Arial", size=12)
            pdf.cell(200, 6, txt=f"Fecha: {record.date.strftime('%d/%m/%Y')}", ln=True)
            pdf.multi_cell(0, 6, txt=f"Informe: {record.report_text}")

            if record.subject_photos:
                pdf.ln(5)
                pdf.cell(200, 6, txt="Fotos del Sujeto:", ln=True)
                x_start = 10
                for photo in record.subject_photos:
                    img_path = os.path.join(current_app.config['UPLOAD_FOLDER'], photo.filename)
                    if os.path.exists(img_path):
                        pdf.image(img_path, x=x_start, y=pdf.get_y(), w=50)
                        x_start += 55
                        if x_start > 150: 
                            x_start = 10
                            pdf.ln(55)
                pdf.ln(60)

            if record.evidence_photos:
                pdf.ln(5)
                pdf.cell(200, 6, txt="Evidencia:", ln=True)
                x_start = 10
                for photo in record.evidence_photos:
                    img_path = os.path.join(current_app.config['UPLOAD_FOLDER'], photo.filename)
                    if os.path.exists(img_path):
                        pdf.image(img_path, x=x_start, y=pdf.get_y(), w=50)
                        x_start += 55
                        if x_start > 150:
                            x_start = 10
                            pdf.ln(55)
                pdf.ln(60)

            pdf.ln(10)
            pdf.line(10, pdf.get_y(), 200, pdf.get_y())
            pdf.ln(10)

    pdf_bytes = pdf.output()
    response = make_response(bytes(pdf_bytes))
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f'attachment; filename=antecedentes_{current_user.dni}.pdf'
    return response

@bp.route('/judicial')
@login_required
def judicial():
    if current_user.badge_id:
        return redirect(url_for('main.official_dashboard'))
    return render_template('judicial.html')

# --- PLANTILLAS ROUTES ---

@bp.route('/official/plantillas/generate_sabes', methods=['POST'])
@login_required
def generate_sabes_report():
    # Only allow officials
    if not current_user.badge_id:
        return redirect(url_for('main.index'))

    # Get data from form
    nombre_agente = request.form.get('nombre_agente')
    fecha = request.form.get('fecha')
    detalles = request.form.get('detalles')
    photo = request.files.get('evidence_photo')
    titulo = request.form.get('titulo')
    directed_to = request.form.get('directed_to')

    # Initialize FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    # Header / Title
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, txt=titulo if titulo else "Reporte SABES", ln=True, align='C')
    pdf.ln(10)

    # Metadata
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 8, txt=f"Fecha: {fecha}", ln=True)
    pdf.cell(0, 8, txt=f"Agente: {nombre_agente}", ln=True)
    if directed_to:
        pdf.cell(0, 8, txt=f"Dirigido a: {directed_to}", ln=True)
    pdf.ln(10)

    # Details
    pdf.set_font("Arial", 'B', 12)
    pdf.cell(0, 10, txt="Detalles:", ln=True)
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 6, txt=detalles)
    pdf.ln(10)

    # Evidence Photo
    if photo and photo.filename:
        filename = secure_filename(photo.filename)
        if not os.path.exists(current_app.config['UPLOAD_FOLDER']):
            os.makedirs(current_app.config['UPLOAD_FOLDER'])

        temp_path = os.path.join(current_app.config['UPLOAD_FOLDER'], filename)
        photo.save(temp_path)

        pdf.ln(5)
        pdf.set_font("Arial", 'B', 12)
        pdf.cell(0, 10, txt="Evidencia Gráfica:", ln=True)

        try:
            pdf.image(temp_path, x=10, w=150)
        except Exception as e:
            pdf.cell(0, 10, txt=f"[Error al adjuntar imagen: {str(e)}]", ln=True)

    # Output
    pdf_bytes = pdf.output()
    response = make_response(bytes(pdf_bytes))
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = f'attachment; filename=Reporte_SABES_{fecha}.pdf'
    return response

@bp.route('/licenses', methods=['GET', 'POST'])
@login_required
def licenses():
    if current_user.badge_id:
        return redirect(url_for('main.official_dashboard'))

    # Precios de licencias personales (Ahora Gratis)
    personal_license_prices = {
        'Pilot': 0,
        'Stripping': 0
    }
    
    license_names = {
        'Pilot': 'Licencia de Piloto',
        'Stripping': 'Licencia de Stripping'
    }

    business_form = BusinessLicenseForm()

    # Check for expired licenses
    expired_count = 0
    for lic in current_user.licenses:
        if lic.status == 'Activa' and lic.expiration_date and lic.expiration_date < datetime.utcnow().date():
            lic.status = 'Vencida'
            expired_count += 1

    if expired_count > 0:
        db.session.commit()
        flash(f'¡Atención! Tienes {expired_count} licencia(s) vencida(s). Renuevalas cuanto antes.', 'warning')

    if request.method == 'POST':
        # --- Lógica de Compra de Licencias Personales ---
        if 'licenses' in request.form:
            selected_licenses = request.form.getlist('licenses')
            if not selected_licenses:
                flash('No seleccionaste ninguna licencia.')
                return redirect(url_for('main.licenses'))

            valid_licenses = []
            for lic in selected_licenses:
                if lic in personal_license_prices:
                    valid_licenses.append(lic)

            # status='Pendiente', expiration_date=None until approved
            for lic in valid_licenses:
                new_license = License(
                    type=license_names.get(lic, lic),
                    status='Pendiente',
                    issue_date=None,
                    expiration_date=None,
                    user_id=current_user.id
                )
                db.session.add(new_license)

            db.session.commit()
            flash(f'Solicitud de licencias enviada. Espera la aprobación de un agente del SABES.')
            return redirect(url_for('main.licenses'))

    return render_template('licenses.html', 
                           personal_prices=personal_license_prices,
                           active_licenses=current_user.licenses,
                           business_form=business_form)

@bp.route('/licenses/business/register', methods=['POST'])
@login_required
def register_business():
    form = BusinessLicenseForm()

    if form.validate_on_submit():
        extra_license_name = None

        b_type = form.business_type.data
        
        # Mapa de nombres
        if b_type == '247':
            extra_license_name = 'Venta Alcohol y Tabaco'
        elif b_type == 'Pharmacy':
            extra_license_name = 'Venta Drogas Farmacas'
        elif b_type == 'Mechanic':
            extra_license_name = 'Reparación de Vehículos'
        elif b_type in ['Restaurant', 'GasStation', 'Club', 'Bar']:
            extra_license_name = 'Venta Alcohol y Tabaco'
        elif b_type == 'UsedCars':
            extra_license_name = 'Venta Vehículos Usados'
        
        # Guardar Foto
        photo_filename = None
        if form.photo.data:
            f = form.photo.data
            photo_filename = secure_filename(f.filename)
            f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], photo_filename))

        # Crear Negocio
        new_business = Business(
            name=form.name.data,
            type=b_type,
            location_x=form.location_x.data,
            location_y=form.location_y.data,
            photo_filename=photo_filename,
            owner_id=current_user.id
        )
        db.session.add(new_business)
        db.session.commit() # Commit para obtener el ID del negocio

        # Crear Licencias vinculadas al negocio
        # Estado inicial Pendiente
        
        # 1. Licencia de Funcionamiento (Obligatoria)
        lic1 = License(
            type='Licencia de Funcionamiento',
            status='Pendiente',
            issue_date=None,
            expiration_date=None,
            user_id=current_user.id,
            business_id=new_business.id
        )
        db.session.add(lic1)

        # 2. Licencia Extra (si aplica)
        if extra_license_name:
            lic2 = License(
                type=extra_license_name,
                status='Pendiente',
                issue_date=None,
                expiration_date=None,
                user_id=current_user.id,
                business_id=new_business.id
            )
            db.session.add(lic2)

        db.session.commit()
        flash(f'Negocio "{form.name.data}" registrado. Licencias pendientes de aprobación por SABES.')
        return redirect(url_for('main.licenses'))
    
    else:
        flash('Error en el formulario. Revisa los datos.')
        return redirect(url_for('main.licenses'))

# --- OFFICIAL ROUTES ---

@bp.route('/official/login', methods=['GET', 'POST'])
def official_login():
    if current_user.is_authenticated:
        if current_user.badge_id:
             return redirect(url_for('main.official_dashboard'))
        return redirect(url_for('main.citizen_dashboard'))

    form = OfficialLoginForm()
    if form.validate_on_submit():
        user = User.query.filter_by(badge_id=form.badge_id.data).first()
        if user is None or not user.check_password(form.password.data):
            flash('Placa ID o contraseña inválidos')
            return redirect(url_for('main.official_login'))

        if user.official_status != 'Aprobado':
             flash('Tu cuenta aún no ha sido aprobada por un líder.')
             return redirect(url_for('main.official_login'))

        login_user(user, remember=form.remember_me.data)
        return redirect(url_for('main.official_dashboard'))

    return render_template('official_login.html', form=form)

@bp.route('/official/register', methods=['GET', 'POST'])
def official_register():
    if current_user.is_authenticated:
        return redirect(url_for('main.official_dashboard'))

    form = OfficialRegistrationForm()
    if form.validate_on_submit():
        citizen = User.query.filter_by(dni=form.dni.data, badge_id=None).first()
        if not citizen:
            flash('Debes estar registrado como ciudadano primero (DNI no encontrado).')
            return redirect(url_for('main.official_register'))

        if not citizen.check_password(form.password.data):
             flash('Contraseña incorrecta. Usa tu contraseña de ciudadano.')
             return redirect(url_for('main.official_register'))

        if User.query.filter_by(badge_id=form.badge_id.data).first():
            flash('Esa Placa ID ya está registrada.')
            return redirect(url_for('main.official_register'))

        photo_file = form.photo.data
        photo_filename = secure_filename(photo_file.filename)

        if not os.path.exists(current_app.config['UPLOAD_FOLDER']):
            os.makedirs(current_app.config['UPLOAD_FOLDER'])

        photo_path = os.path.join(current_app.config['UPLOAD_FOLDER'], photo_filename)
        photo_file.save(photo_path)

        citizen.badge_id = form.badge_id.data
        citizen.department = form.department.data
        citizen.selfie_filename = photo_filename
        citizen.official_status = 'Pendiente'
        citizen.official_rank = 'Miembro'

        # citizen.set_password(form.password.data) # Password verified above

        db.session.commit()

        flash('Solicitud enviada. Espera a que un líder apruebe tu cuenta.')
        return redirect(url_for('main.official_login'))

    return render_template('official_register.html', form=form)

@bp.route('/official/dashboard')
@login_required
def official_dashboard():
    if not current_user.badge_id:
        return redirect(url_for('main.citizen_dashboard'))

    pending_users = []
    if current_user.department == 'Gobierno' and current_user.official_rank == 'Lider':
         pending_users = User.query.filter_by(official_status='Pendiente').all()
    elif current_user.official_rank == 'Lider':
        pending_users = User.query.filter_by(department=current_user.department, official_status='Pendiente').all()

    return render_template('official_dashboard.html', pending_users=pending_users)

@bp.route('/official/action/<int:user_id>/<action>', methods=['POST'])
@login_required
def official_action(user_id, action):
    if not current_user.badge_id or current_user.official_rank != 'Lider':
        return redirect(url_for('main.citizen_dashboard'))

    target_user = User.query.get_or_404(user_id)

    # MODIFICADO: Permitir si es del mismo departamento O si el usuario actual es de 'Gobierno'
    if target_user.department != current_user.department and current_user.department != 'Gobierno':
        flash('No tienes permiso para gestionar este usuario.')
        return redirect(url_for('main.official_dashboard'))

    if action == 'approve':
        target_user.official_status = 'Aprobado'
        flash(f'Usuario {target_user.first_name} {target_user.last_name} aprobado.')
    elif action == 'deny':
        db.session.delete(target_user)
        flash(f'Usuario {target_user.first_name} {target_user.last_name} denegado y eliminado.')

    db.session.commit()
    return redirect(url_for('main.official_dashboard'))

@bp.route('/official/licenses/pending')
@login_required
def official_licenses_pending():
    if not current_user.badge_id or current_user.department != 'SABES':
        flash('Acceso denegado. Solo personal de SABES.')
        return redirect(url_for('main.official_dashboard'))

    pending_licenses = License.query.filter_by(status='Pendiente').all()
    return render_template('manage_licenses.html', licenses=pending_licenses)

@bp.route('/official/licenses/action/<int:license_id>/<action>', methods=['POST'])
@login_required
def official_license_action(license_id, action):
    if not current_user.badge_id or current_user.department != 'SABES':
        flash('Acceso denegado.')
        return redirect(url_for('main.official_dashboard'))

    lic = License.query.get_or_404(license_id)

    if action == 'approve':
        lic.status = 'Activa'
        lic.issue_date = datetime.utcnow().date()
        lic.expiration_date = datetime.utcnow().date() + timedelta(days=30)
        flash(f'Licencia {lic.type} aprobada para {lic.holder.first_name} {lic.holder.last_name}. Expira en 30 días.')

        notify_discord_bot(lic.holder, f"✅ **Licencia Aprobada**\nTu licencia '{lic.type}' ha sido aprobada y es válida hasta el {lic.expiration_date}.")

    elif action == 'reject':
        lic.status = 'Rechazada'
        flash(f'Licencia {lic.type} rechazada.')
        notify_discord_bot(lic.holder, f"❌ **Licencia Rechazada**\nTu solicitud para la licencia '{lic.type}' ha sido rechazada. Contacta a SABES para más información.")

    db.session.commit()
    return redirect(url_for('main.official_licenses_pending'))

# --- GOVERNMENT DASHBOARD & PAYROLL ---

@bp.route('/government/dashboard')
@login_required
def government_dashboard():
    if current_user.department != 'Gobierno':
        return redirect(url_for('main.official_dashboard'))

    create_leader_form = CreateLeaderForm()

    # Obtener lista de todos los usuarios ciudadanos (no funcionarios)
    all_users = User.query.filter(User.badge_id == None).all()

    return render_template('government_dashboard.html',
                           create_leader_form=create_leader_form,
                           all_users=all_users)

@bp.route('/government/create_leader', methods=['POST'])
@login_required
def government_create_leader():
    if current_user.department != 'Gobierno':
        return redirect(url_for('main.official_dashboard'))

    form = CreateLeaderForm()
    if form.validate_on_submit():
        citizen = User.query.filter_by(dni=form.dni.data, badge_id=None).first()
        if not citizen:
            flash('El ciudadano no existe (DNI no encontrado). Debe registrarse primero.')
            return redirect(url_for('main.government_dashboard'))

        if User.query.filter_by(badge_id=form.badge_id.data).first():
            flash('Esa Placa ID ya está registrada.')
            return redirect(url_for('main.government_dashboard'))

        user = User(
            first_name=form.first_name.data,
            last_name=form.last_name.data,
            dni=form.dni.data,
            badge_id=form.badge_id.data,
            department=form.department.data,
            official_status='Aprobado',
            official_rank='Lider',
            selfie_filename='default.jpg',
            dni_photo_filename='default.jpg'
        )
        user.set_password(form.password.data)

        db.session.add(user)
        db.session.commit()

        flash(f'Líder de {form.department.data} creado con éxito.')
    else:
        for field, errors in form.errors.items():
            for error in errors:
                flash(f"Error en {getattr(form, field).label.text}: {error}")

    return redirect(url_for('main.government_dashboard'))

@bp.route('/government/users', methods=['GET'])
@login_required
def government_users():
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.official_dashboard'))

    users = User.query.all()
    return render_template('government_users.html', users=users)

@bp.route('/government/users/bulk_action', methods=['POST'])
@login_required
def government_users_bulk_action():
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.official_dashboard'))

    action = request.form.get('action')
    user_ids = request.form.getlist('user_ids')

    if not user_ids:
        flash('No seleccionaste ningún usuario.')
        return redirect(url_for('main.government_users'))

    count = 0

    if action == 'unlink_discord':
        for uid in user_ids:
            u = User.query.get(uid)
            if u:
                u.discord_id = None
                count += 1
        db.session.commit()
        flash(f'Discord desvinculado de {count} usuarios.')

    elif action == 'delete':
        for uid in user_ids:
            u = User.query.get(uid)
            if u:
                if u.id == current_user.id:
                    continue # Skip self
                _perform_user_deletion(u)
                count += 1
        db.session.commit()
        flash(f'{count} usuarios eliminados permanentemente.')

    return redirect(url_for('main.government_users'))

@bp.route('/official/kick_member/<int:user_id>', methods=['POST'])
@login_required
def kick_member(user_id):
    if not current_user.badge_id or current_user.official_rank != 'Lider':
        flash('No tienes permiso para realizar esta acción.', 'danger')
        return redirect(url_for('main.official_dashboard'))

    target_user = User.query.get_or_404(user_id)

    # Validar permisos: Mismo departamento O Gobierno
    if target_user.department != current_user.department and current_user.department != 'Gobierno':
        flash('No puedes expulsar a miembros de otro departamento.', 'danger')
        return redirect(url_for('main.official_dashboard'))
    
    # Evitar auto-expulsión accidental (opcional, pero recomendado)
    if target_user.id == current_user.id:
        flash('No puedes expulsarte a ti mismo.', 'warning')
        return redirect(url_for('main.official_dashboard'))

    # Lógica de expulsión
    target_user.official_status = 'Suspendido'
    target_user.badge_id = None # Revocar acceso oficial
    # Opcional: target_user.department = None (Si quieres que dejen de pertenecer al dpto totalmente)
    
    db.session.commit()
    
    flash(f'Funcionario {target_user.first_name} {target_user.last_name} ha sido expulsado del departamento.', 'success')
    return redirect(url_for('main.official_dashboard'))

# --- CITIZEN DATABASE ROUTES ---

@bp.route('/official/database', methods=['GET'])
@login_required
def official_database():
    if not current_user.badge_id:
        return redirect(url_for('main.citizen_dashboard'))

    form = SearchUserForm(request.args)
    users = []
    if form.query.data:
        query = form.query.data
        users = User.query.filter(
            (
                User.first_name.contains(query) |
                User.last_name.contains(query) |
                User.dni.contains(query)
            )
        ).all()

    return render_template('official_database.html', form=form, users=users)

@bp.route('/official/citizen/<int:user_id>')
@login_required
def citizen_profile(user_id):
    if not current_user.badge_id:
        return redirect(url_for('main.citizen_dashboard'))

    citizen = User.query.get_or_404(user_id)

    can_edit_reports = current_user.department in ['SABES', 'Gobierno']

    comment_form = CommentForm()
    fine_form = TrafficFineForm()
    criminal_form = CriminalRecordForm()
    
    # Nuevos formularios
    edit_info_form = EditCitizenForm(obj=citizen)
    edit_photos_form = EditCitizenPhotoForm()
    # NUEVO: Formulario cambio de contraseña para admin
    change_password_form = ChangePasswordForm()

    return render_template('citizen_profile.html', citizen=citizen,
                           can_edit=can_edit_reports,
                           comment_form=comment_form, fine_form=fine_form,
                           criminal_form=criminal_form,
                           edit_info_form=edit_info_form, edit_photos_form=edit_photos_form,
                           change_password_form=change_password_form)

@bp.route('/official/citizen/<int:user_id>/add_comment', methods=['POST'])
@login_required
def add_comment(user_id):
    if not current_user.badge_id:
        return redirect(url_for('main.citizen_dashboard'))

    form = CommentForm()
    if form.validate_on_submit():
        comment = Comment(
            content=form.content.data,
            user_id=user_id,
            author_id=current_user.id
        )
        db.session.add(comment)
        db.session.commit()
        flash('Comentario agregado.')
    else:
        flash('Error al agregar comentario.')

    return redirect(url_for('main.citizen_profile', user_id=user_id))


# --- SA FINDER ROUTES ---

@bp.route('/official/safinder', methods=['GET'])
def safinder():
    # Public access allowed for viewing/search
    # Uploads restricted to officials via template logic and separate upload route

    query = request.args.get('q', '')
    results = []

    if query:
        search = f"%{query}%"
        results = DocModel.query.filter(
            (DocModel.title.ilike(search)) |
            (DocModel.text_content.ilike(search))
        ).order_by(DocModel.created_at.desc()).all()

    recent_docs = DocModel.query.order_by(DocModel.created_at.desc()).limit(5).all()

    return render_template('safinder.html', results=results, recent_docs=recent_docs)

@bp.route('/official/safinder/upload', methods=['POST'])
@login_required
def safinder_upload():
    if not current_user.badge_id:
        return redirect(url_for('main.citizen_dashboard'))

    if 'file' not in request.files:
        flash('No se seleccionó archivo.')
        return redirect(url_for('main.safinder'))

    file = request.files['file']
    title = request.form.get('title')

    if file.filename == '':
        flash('Nombre de archivo inválido.')
        return redirect(url_for('main.safinder'))

    if file and file.filename.lower().endswith('.pdf'):
        import uuid
        filename = f"{uuid.uuid4().hex}_{secure_filename(file.filename)}"
        docs_folder = os.path.join(current_app.config['UPLOAD_FOLDER'], 'docs')

        if not os.path.exists(docs_folder):
            os.makedirs(docs_folder)

        file_path = os.path.join(docs_folder, filename)
        file.save(file_path)

        # Extract Text
        try:
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
        except Exception as e:
            print(f"Error parsing PDF: {e}")
            text = "Error leyendo contenido."

        new_doc = DocModel(
            title=title,
            filename=filename,
            text_content=text,
            uploader_id=current_user.id
        )
        db.session.add(new_doc)
        db.session.commit()

        flash('Documento subido e indexado correctamente.')
    else:
        flash('Solo se permiten archivos PDF.')

    return redirect(url_for('main.safinder'))

@bp.route('/official/citizen/<int:user_id>/add_traffic_fine', methods=['POST'])
@login_required
def add_traffic_fine(user_id):
    if not current_user.badge_id:
        return redirect(url_for('main.citizen_dashboard'))

    form = TrafficFineForm()
    citizen = User.query.get_or_404(user_id)
    if form.validate_on_submit():
        fine = TrafficFine(
            reason=form.reason.data,
            user_id=user_id,
            author_id=current_user.id
        )
        db.session.add(fine)
        db.session.commit()
        
        notify_discord_bot(citizen, f"🚨 **Has recibido una Multa**\nRazón: {form.reason.data}\nAgente: {current_user.first_name} {current_user.last_name} ({current_user.department})")

        flash(f'Multa impuesta.')
    else:
        flash('Error al imponer multa.')

    return redirect(url_for('main.citizen_profile', user_id=user_id))

@bp.route('/official/citizen/<int:user_id>/add_criminal_record', methods=['POST'])
@login_required
def add_criminal_record(user_id):
    if not current_user.badge_id:
        return redirect(url_for('main.citizen_dashboard'))

    if current_user.department not in ['SABES', 'Gobierno']:
         flash('No tienes permiso para agregar antecedentes penales.')
         return redirect(url_for('main.citizen_profile', user_id=user_id))
    
    citizen = User.query.get_or_404(user_id)
    form = CriminalRecordForm()
    if form.validate_on_submit():
        record = CriminalRecord(
            date=form.date.data,
            crime=form.crime.data,
            penal_code=form.penal_code.data,
            report_text=form.report_text.data,
            user_id=user_id,
            author_id=current_user.id
        )
        db.session.add(record)
        db.session.commit() # Commit inicial para tener el ID del record

        if not os.path.exists(current_app.config['UPLOAD_FOLDER']):
            os.makedirs(current_app.config['UPLOAD_FOLDER'])

        for file in form.subject_photos.data:
            if file and file.filename != '':
                filename = secure_filename(file.filename)
                file.save(os.path.join(current_app.config['UPLOAD_FOLDER'], filename))
                photo = CriminalRecordSubjectPhoto(filename=filename, record_id=record.id)
                db.session.add(photo)

        for file in form.evidence_photos.data:
             if file and file.filename != '':
                filename = secure_filename(file.filename)
                file.save(os.path.join(current_app.config['UPLOAD_FOLDER'], filename))
                photo = CriminalRecordEvidencePhoto(filename=filename, record_id=record.id)
                db.session.add(photo)

        db.session.commit()
        
        notify_discord_bot(citizen, f"⚖️ **Nuevo Antecedente Penal**\nDelito: {form.crime.data}\nCódigo Penal: {form.penal_code.data}\nAgente: {current_user.first_name} {current_user.last_name}")
        
        flash('Antecedente penal registrado.')
    else:
        flash('Error en el formulario.')

    return redirect(url_for('main.citizen_profile', user_id=user_id))

# --- NUEVAS RUTAS ADMINISTRATIVAS DE GOBIERNO ---

@bp.route('/official/citizen/<int:user_id>/edit_info', methods=['POST'])
@login_required
def edit_citizen_info(user_id):
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.citizen_profile', user_id=user_id))
    
    user = User.query.get_or_404(user_id)
    form = EditCitizenForm()
    
    if form.validate_on_submit():
        user.first_name = form.first_name.data
        user.last_name = form.last_name.data
        user.dni = form.dni.data
        db.session.commit()
        flash('Información personal actualizada exitosamente.')
    else:
        flash('Error al actualizar información.')
    
    return redirect(url_for('main.citizen_profile', user_id=user_id))

@bp.route('/official/citizen/<int:user_id>/update_photos', methods=['POST'])
@login_required
def update_citizen_photos(user_id):
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.citizen_profile', user_id=user_id))

    user = User.query.get_or_404(user_id)
    form = EditCitizenPhotoForm()

    if form.validate_on_submit():
        if form.selfie.data:
            f = form.selfie.data
            filename = secure_filename(f.filename)
            f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], filename))
            user.selfie_filename = filename
        
        if form.dni_photo.data:
            f = form.dni_photo.data
            filename = secure_filename(f.filename)
            f.save(os.path.join(current_app.config['UPLOAD_FOLDER'], filename))
            user.dni_photo_filename = filename
            
        db.session.commit()
        flash('Fotos actualizadas exitosamente.')
    else:
        flash('Error al subir fotos. Verifica el formato.')

    return redirect(url_for('main.citizen_profile', user_id=user_id))

@bp.route('/official/citizen/<int:user_id>/unlink_discord', methods=['POST'])
@login_required
def unlink_discord(user_id):
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.citizen_profile', user_id=user_id))
        
    user = User.query.get_or_404(user_id)
    user.discord_id = None
    db.session.commit()
    flash('Discord desvinculado exitosamente.')
    return redirect(url_for('main.citizen_profile', user_id=user_id))

@bp.route('/official/citizen/<int:user_id>/clear_records', methods=['POST'])
@login_required
def clear_criminal_records(user_id):
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.citizen_profile', user_id=user_id))
        
    user = User.query.get_or_404(user_id)
    # Borrar todos los antecedentes de este usuario específico
    count = CriminalRecord.query.filter_by(user_id=user.id).delete()
    db.session.commit()
    flash(f'Se han borrado {count} antecedentes penales de este ciudadano.')
    return redirect(url_for('main.citizen_profile', user_id=user_id))

# RUTA PARA CAMBIAR CONTRASEÑA (Reemplaza a reset_account)
@bp.route('/official/citizen/<int:user_id>/change_password', methods=['POST'])
@login_required
def change_citizen_password(user_id):
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.citizen_profile', user_id=user_id))
        
    user = User.query.get_or_404(user_id)
    form = ChangePasswordForm()
    
    if form.validate_on_submit():
        user.set_password(form.new_password.data)
        db.session.commit()
        flash(f'Contraseña de {user.first_name} cambiada exitosamente.')
    else:
        flash('Error al cambiar contraseña.')
        
    return redirect(url_for('main.citizen_profile', user_id=user_id))

@bp.route('/official/citizen/<int:user_id>/delete_account', methods=['POST'])
@login_required
def delete_citizen_account(user_id):
    if current_user.department != 'Gobierno':
        flash('Acceso denegado.')
        return redirect(url_for('main.citizen_profile', user_id=user_id))

    user = User.query.get_or_404(user_id)

    _perform_user_deletion(user)
    db.session.commit()

    flash(f'Usuario {user.first_name} {user.last_name} eliminado permanentemente.')
    return redirect(url_for('main.official_dashboard'))

@bp.route('/licenses/business/<int:business_id>/transfer', methods=['POST'])
@login_required
def transfer_business(business_id):
    business = Business.query.get_or_404(business_id)
    if business.owner_id != current_user.id:
        flash('No tienes permiso para transferir este negocio.', 'danger')
        return redirect(url_for('main.licenses'))

    new_owner_dni = request.form.get('new_owner_dni')
    new_owner = User.query.filter_by(dni=new_owner_dni).first()

    if not new_owner:
        flash('El usuario con ese DNI no existe.', 'danger')
        return redirect(url_for('main.licenses'))

    if new_owner.id == current_user.id:
        flash('No puedes transferirte el negocio a ti mismo.', 'warning')
        return redirect(url_for('main.licenses'))

    # Transfer ownership
    business.owner_id = new_owner.id

    # Transfer associated licenses? Usually licenses are tied to business, so owner change is enough if logic uses business.owner
    # Our License model has user_id. We need to update user_id on licenses too.
    for lic in business.licenses:
        lic.user_id = new_owner.id

    db.session.commit()

    notify_discord_bot(new_owner, f"🏢 **Nuevo Negocio Recibido**\n{current_user.first_name} {current_user.last_name} te ha transferido el negocio '{business.name}'.")
    notify_discord_bot(current_user, f"🏢 **Negocio Transferido**\nHas transferido '{business.name}' a {new_owner.first_name} {new_owner.last_name}.")

    flash(f'Negocio "{business.name}" transferido exitosamente a {new_owner.first_name} {new_owner.last_name}.', 'success')
    return redirect(url_for('main.licenses'))

@bp.route('/licenses/business/<int:business_id>/pay_fine/<int:fine_id>', methods=['POST'])
@login_required
def pay_business_fine(business_id, fine_id):
    business = Business.query.get_or_404(business_id)
    if business.owner_id != current_user.id:
        flash('Acceso denegado.', 'danger')
        return redirect(url_for('main.licenses'))

    fine = BusinessFine.query.get_or_404(fine_id)
    if fine.business_id != business.id:
        flash('Multa no corresponde al negocio.', 'danger')
        return redirect(url_for('main.licenses'))

    if fine.status == 'Pagada':
        flash('Esta multa ya está pagada.', 'info')
    else:
        fine.status = 'Pagada'
        db.session.commit()
        flash('Multa pagada exitosamente.', 'success')

    return redirect(url_for('main.licenses'))

@bp.route('/licenses/business/<int:business_id>/renew_license/<int:license_id>', methods=['POST'])
@login_required
def renew_business_license(business_id, license_id):
    business = Business.query.get_or_404(business_id)
    if business.owner_id != current_user.id:
        flash('Acceso denegado.', 'danger')
        return redirect(url_for('main.licenses'))

    lic = License.query.get_or_404(license_id)
    if lic.business_id != business.id:
        flash('Licencia no corresponde al negocio.', 'danger')
        return redirect(url_for('main.licenses'))

    # Logic: Set to Pendiente for approval
    lic.status = 'Pendiente'
    db.session.commit()

    flash(f'Solicitud de renovación para "{lic.type}" enviada. Espera aprobación.', 'success')
    return redirect(url_for('main.licenses'))

# --- OFFICIAL BUSINESS ROUTES ---

@bp.route('/official/businesses')
@login_required
def official_businesses():
    allowed_departments = ['SABES', 'Gobierno']
    if current_user.department not in allowed_departments:
        flash('Acceso denegado.', 'danger')
        return redirect(url_for('main.official_dashboard'))

    query = request.args.get('q', '')
    if query:
        search = f"%{query}%"
        businesses = Business.query.filter(Business.name.ilike(search)).all()
    else:
        businesses = Business.query.all()

    return render_template('official_businesses.html', businesses=businesses)

@bp.route('/official/business/<int:business_id>/fine', methods=['POST'])
@login_required
def official_business_fine(business_id):
    allowed_departments = ['SABES', 'Gobierno']
    if current_user.department not in allowed_departments:
        flash('Acceso denegado.', 'danger')
        return redirect(url_for('main.official_dashboard'))

    business = Business.query.get_or_404(business_id)
    reason = request.form.get('reason')

    if not reason:
        flash('Debes especificar una razón.', 'warning')
        return redirect(url_for('main.official_businesses'))

    fine = BusinessFine(
        reason=reason,
        business_id=business.id,
        author_id=current_user.id,
        status='Pendiente'
    )
    db.session.add(fine)
    db.session.commit()

    notify_discord_bot(business.owner, f"🚨 **Multa a Negocio**\nTu negocio '{business.name}' ha recibido una multa.\nRazón: {reason}\nAgente: {current_user.first_name} {current_user.last_name}")

    flash(f'Multa aplicada a "{business.name}".', 'success')
    return redirect(url_for('main.official_businesses'))

@bp.route('/official/business/<int:business_id>/approve_registration', methods=['POST'])
@login_required
def official_approve_business_registration(business_id):
    allowed_departments = ['SABES', 'Gobierno']
    if current_user.department not in allowed_departments:
        flash('Acceso denegado.', 'danger')
        return redirect(url_for('main.official_dashboard'))

    business = Business.query.get_or_404(business_id)
    business.status = 'Aprobado'

    # Also approve the initial "Licencia de Funcionamiento" if present and pending?
    # Usually registration approval implies functionality approval.
    for lic in business.licenses:
        if lic.type == 'Licencia de Funcionamiento' and lic.status == 'Pendiente':
            lic.status = 'Activa'
            lic.issue_date = datetime.utcnow().date()
            lic.expiration_date = datetime.utcnow().date() + timedelta(days=30)

    db.session.commit()

    notify_discord_bot(business.owner, f"✅ **Negocio Aprobado**\nTu negocio '{business.name}' ha sido registrado y aprobado exitosamente.")

    flash(f'Negocio "{business.name}" aprobado.', 'success')
    return redirect(url_for('main.official_businesses'))

@bp.route('/official/business/<int:business_id>/reject_registration', methods=['POST'])
@login_required
def official_reject_business_registration(business_id):
    allowed_departments = ['SABES', 'Gobierno']
    if current_user.department not in allowed_departments:
        flash('Acceso denegado.', 'danger')
        return redirect(url_for('main.official_dashboard'))

    business = Business.query.get_or_404(business_id)
    # Delete business or set to Rejected? User said "Once accepted... it appears".
    # Usually rejection means you have to apply again or fix it.
    # Deleting cleans up.

    notify_discord_bot(business.owner, f"❌ **Registro Rechazado**\nLa solicitud para tu negocio '{business.name}' ha sido rechazada.")

    db.session.delete(business)
    db.session.commit()

    flash(f'Solicitud de negocio rechazada y eliminada.', 'warning')
    return redirect(url_for('main.official_businesses'))
